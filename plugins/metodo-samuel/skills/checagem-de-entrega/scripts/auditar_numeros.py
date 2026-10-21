"""Auditoria rápida de números em planilha (.xlsx/.xlsm) ou CSV.

Aponta o que a checagem-de-entrega rejeita:
  - linha ou coluna de total que não bate com a soma das partes
  - números redondos demais (muitos múltiplos de 100 numa coluna)
  - 3 ou mais valores idênticos em sequência (cara de "marretado")
  - salto entre períodos consecutivos acima do limite (padrão 25%)
  - fórmulas sem valor calculado (arquivo salvo sem recalcular)

Uso:
  python auditar_numeros.py arquivo.xlsx [--aba NOME] [--cabecalho 0] [--salto 0.25] [--json]

Não altera o arquivo. Só lê e relata.
"""
import argparse
import json
import sys

import pandas as pd

TOTAL_TOKENS = ("total", "subtotal", "soma", "geral")
TOL_REL = 0.0001  # 0,01%: só absorve arredondamento de casas decimais


def eh_total(texto):
    return isinstance(texto, str) and any(t in texto.strip().lower() for t in TOTAL_TOKENS)


def carregar(caminho, aba, cabecalho):
    if caminho.lower().endswith((".csv", ".txt")):
        df = pd.read_csv(caminho, header=cabecalho, sep=None, engine="python", decimal=",")
        return {"csv": df}, []
    avisos = []
    try:
        import openpyxl

        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        wb_f = openpyxl.load_workbook(caminho, data_only=False, read_only=True)
        for ws in wb.worksheets:
            if aba and ws.title != aba:
                continue
            sem_valor = 0
            for row_v, row_f in zip(ws.iter_rows(values_only=True), wb_f[ws.title].iter_rows(values_only=True)):
                for v, f in zip(row_v, row_f):
                    if v is None and isinstance(f, str) and f.startswith("="):
                        sem_valor += 1
            if sem_valor:
                avisos.append(
                    f"[{ws.title}] {sem_valor} fórmula(s) sem valor calculado: abrir e salvar no Excel "
                    "(ou recalcular com LibreOffice) antes de auditar."
                )
    except Exception as e:  # arquivo protegido, formato antigo etc.
        avisos.append(f"Não consegui checar fórmulas: {e}")
    abas = pd.read_excel(caminho, sheet_name=aba if aba else None, header=cabecalho)
    if isinstance(abas, pd.DataFrame):
        abas = {aba: abas}
    return abas, avisos


CHAVE_TOKENS = ("mes", "mês", "ano", "data", "periodo", "período", "codigo", "código", "cod_", "id", "sku", "cnpj", "numero_", "número")


def eh_chave(nome, s):
    """Coluna de período ou identificador (ex.: numero_ano_mes, id): não é dado a auditar."""
    n = str(nome).strip().lower()
    if any(n == t or n.startswith(t) or n.endswith(t) for t in CHAVE_TOKENS):
        return True
    s = s.dropna()
    return len(s) >= 3 and (s % 1 == 0).all() and s.is_monotonic_increasing and s.nunique() == len(s) and s.diff().dropna().max() <= 100


def colunas_numericas(df):
    cols = []
    for c in df.columns:
        s = pd.to_numeric(df[c], errors="coerce")
        if s.notna().sum() >= max(3, 0.6 * df[c].notna().sum()) and not eh_chave(c, s):
            cols.append(c)
    return cols


def auditar_aba(nome, df, limite_salto):
    achados = []
    df = df.dropna(how="all").dropna(axis=1, how="all")
    if df.empty:
        return achados
    num = colunas_numericas(df)
    if not num:
        return achados
    rotulo = df[df.columns[0]] if df.columns[0] not in num else pd.Series([None] * len(df), index=df.index)
    linhas_total = [i for i in df.index if eh_total(rotulo.get(i))]
    cols_total = [c for c in num if eh_total(str(c))]
    cols_dado = [c for c in num if c not in cols_total]
    vals = df[num].apply(pd.to_numeric, errors="coerce")

    # 1. Linha de total = soma das linhas acima (desde o total anterior)
    bloco = []
    for i in df.index:
        if i not in linhas_total:
            bloco.append(i)
            continue
        for c in num:
            esperado = vals.loc[bloco, c].sum()
            obtido = vals.at[i, c]
            if bloco and pd.notna(obtido) and abs(obtido - esperado) > max(0.01, TOL_REL * abs(esperado)):
                achados.append(("soma", f"[{nome}] linha '{rotulo[i]}', coluna '{c}': total {obtido:,.2f} x soma das partes {esperado:,.2f}"))
        bloco = []

    # 2. Coluna de total = soma das colunas de dado na mesma linha
    for ct in cols_total:
        for i in df.index:
            esperado = vals.loc[i, cols_dado].sum()
            obtido = vals.at[i, ct]
            if pd.notna(obtido) and vals.loc[i, cols_dado].notna().any() and abs(obtido - esperado) > max(0.01, TOL_REL * abs(esperado)):
                achados.append(("soma", f"[{nome}] linha {i + 2}, coluna '{ct}': {obtido:,.2f} x soma da linha {esperado:,.2f}"))

    base = vals.drop(index=linhas_total, errors="ignore")[cols_dado]

    # 3. Números redondos
    for c in cols_dado:
        s = base[c].dropna()
        s = s[s.abs() >= 100]
        if len(s) >= 4:
            redondos = (s % 100 == 0).mean()
            if redondos >= 0.3:
                achados.append(("redondo", f"[{nome}] coluna '{c}': {redondos:.0%} dos valores ≥100 são múltiplos de 100"))

    # 4. Valores idênticos em sequência (na coluna e na linha)
    def sequencias(serie, onde):
        s = serie.dropna()
        run, ant = 1, None
        for v in s:
            run = run + 1 if v == ant and v != 0 else 1
            if run == 3:
                achados.append(("repetido", f"[{nome}] {onde}: valor {v:,.2f} repetido 3+ vezes em sequência"))
            ant = v

    for c in cols_dado:
        sequencias(base[c], f"coluna '{c}'")
    if len(cols_dado) >= 4:
        for i in base.index:
            sequencias(base.loc[i], f"linha {i + 2}")

    # 5. Saltos entre períodos consecutivos (linha com 4+ períodos = série na horizontal)
    def saltos(serie, onde):
        s = serie.dropna()
        for (k0, a), (k1, b) in zip(s.items(), list(s.items())[1:]):
            if a and abs(a) > 0 and abs(b / a - 1) > limite_salto:
                achados.append(("salto", f"[{nome}] {onde}: {k0} → {k1} variou {b / a - 1:+.0%}"))

    if len(cols_dado) >= 4:
        for i in base.index:
            saltos(base.loc[i], f"linha {i + 2} ('{rotulo.get(i)}')")
    else:
        for c in cols_dado:
            saltos(base[c], f"coluna '{c}'")
    return achados


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("arquivo")
    ap.add_argument("--aba")
    ap.add_argument("--cabecalho", type=int, default=0, help="linha do cabeçalho (0 = primeira)")
    ap.add_argument("--salto", type=float, default=0.25, help="variação máxima entre períodos (0.25 = 25%%)")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()

    abas, avisos = carregar(a.arquivo, a.aba, a.cabecalho)
    achados = []
    for nome, df in abas.items():
        achados += auditar_aba(nome, df, a.salto)

    if a.json:
        print(json.dumps({"avisos": avisos, "achados": [{"tipo": t, "detalhe": d} for t, d in achados]}, ensure_ascii=False, indent=1))
        return
    for av in avisos:
        print(f"AVISO: {av}")
    if not achados:
        print("Nenhum achado: somas fecham, sem redondos em excesso, sem repetição ou salto acima do limite.")
        return
    contagem = pd.Series([t for t, _ in achados]).value_counts()
    print("Resumo: " + ", ".join(f"{k}={v}" for k, v in contagem.items()))
    for tipo in ["soma", "repetido", "redondo", "salto"]:
        itens = [d for t, d in achados if t == tipo]
        if itens:
            print(f"\n## {tipo} ({len(itens)})")
            for d in itens[:30]:
                print(f"- {d}")
            if len(itens) > 30:
                print(f"- ... mais {len(itens) - 30}")
    sys.exit(1 if any(t == "soma" for t, _ in achados) else 0)


if __name__ == "__main__":
    main()
